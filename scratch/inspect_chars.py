
import docx

def inspect_exact_chars(docx_path):
    doc = docx.Document(docx_path)
    # Target Cell(4, 5) which has '91.504.195'
    table = doc.tables[0]
    cell = table.cell(4, 5)
    text = cell.text
    print(f"Cell(4, 5) text: '{text}'")
    print(f"Chars: {[ord(c) for c in text]}")
    
    # Target Cell(12, 2) which has ' 4.882.149.339 '
    cell2 = table.cell(12, 2)
    text2 = cell2.text
    print(f"Cell(12, 2) text: '{text2}'")
    print(f"Chars: {[ord(c) for c in text2]}")

if __name__ == "__main__":
    inspect_exact_chars("number.docx")
